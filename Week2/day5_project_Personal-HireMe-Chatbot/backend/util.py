import os
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
import json
from pypdf import PdfReader
from docx import Document

load_dotenv()
my_api_key=os.getenv("groq_api")

if not my_api_key:
    raise ValueError("API key not found")

client=Groq(api_key=my_api_key)
model = "openai/gpt-oss-safeguard-20b"

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()

class ChatRequest(BaseModel):
    question : str


def parse_resume(resume_text):
    print("LLM called")
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


RESUME_PDF = Path("resume1.pdf")
RESUME_JSON = Path("resume1.json")
def get_resume():

    # If JSON already exists, use it
    if RESUME_JSON.exists():

        with open(RESUME_JSON, "r", encoding="utf-8") as f:
            data = json.load(f)

        return Resume(**data)

    # Otherwise parse the PDF using LLM
    text = read_resume(RESUME_PDF)
    resume = parse_resume(text)

    # Save parsed resume
    with open(RESUME_JSON, "w", encoding="utf-8") as f:
        json.dump(resume.model_dump(), f, indent=2)

    return resume


def ask_candidate(question, resume):
    system_prompt = f"""
        you are an ai assistant representing a job candidate.
        Below is everything you know about the candidate.

        {resume.model_dump_json(indent = 2)}

        Rules:
        1.Answer only using this information
        2.Never hallucinate.
        3.If information is unavailable,
        say "I don't have enough information to answer that."
        4.Be professional
        5.Answer as if HR is interviewing this candidate
        6.Don't give answers in tabular format
        7.When someone asks about the project give the links to github also
    """

    response = client.chat.completions.create(
        model = model,

        messages = [
            {
                "role" : "system",
                "content" : system_prompt
            },

            {
                "role" : "user",
                "content" : question
            }
        ], 
        stream = True
    )
    
    return response